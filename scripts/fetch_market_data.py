"""Download the monthly market series and write data/market/*.csv plus manifest.json.

Sources (decision D3, revised in W14):

- BCRP (Banco Central de Reserva del Perú) statistics, monthly series in soles. The
  BCRP JSON API is protected by an anti-bot challenge, but the public HTML view runs
  that challenge in a real browser, so each series is loaded once with headless
  Chromium (``--dump-dom``) and its table is parsed here.
- Yahoo Finance (optional fallback for stocks, ``--stocks-source yahoo``): the
  ``EPU`` ETF converted to soles with ``PEN=X`` via ``yfinance`` (requirements-data.txt).
- SBS (Superintendencia de Banca, Seguros y AFP), Boletín SPP file B-220932 "Valor
  Cuota diario por Tipo de Fondo de Pensiones y AFP" (W15): the mixed-funds category
  uses the Fondo Tipo 2 (mixed fund by regulation). The latest monthly workbook holds
  the whole daily history; month-end valor cuota levels of Integra, Prima, Profuturo
  (and Habitat, for transparency) are written, plus a chained index of the
  equal-weight average of the three AFPs' monthly returns (what the estimator reads).
  Needs ``openpyxl`` (requirements-data.txt). If the SBS cannot be reached, the
  committed snapshot is kept.

BCRP CSV values are written exactly as published (index level or % annual rate);
months marked "n.d." or empty are skipped. The current, still-open month is dropped.
This is the only place with network access; the app reads the local files only.

Usage: python scripts/fetch_market_data.py [--out data/market] [--start 2010-1]
                                           [--stocks-source bcrp|yahoo]
                                           [--only all|bcrp|sbs]
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import json
import shutil
import subprocess
import urllib.error
import urllib.request
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
BCRP_URL = "https://estadisticas.bcrp.gob.pe/estadisticas/series/mensuales/resultados/{code}/html/{start}/{end}"
CHROMIUM_CANDIDATES = ("chromium", "chromium-browser", "google-chrome", "google-chrome-stable", "chrome")
MAX_ATTEMPTS = 2  # be gentle with BCRP: one request per series, at most one retry
SPANISH_MONTHS = {
    "ene": 1, "feb": 2, "mar": 3, "abr": 4, "may": 5, "jun": 6,
    "jul": 7, "ago": 8, "sep": 9, "set": 9, "oct": 10, "nov": 11, "dic": 12,
}
MISSING_MARKERS = {"", "n.d.", "nd", "n.d", "-", "—"}
MIN_MONTHS = 24
CATEGORIES = ("stocks", "mixed", "debt", "bonds", "term")


@dataclass(frozen=True)
class BcrpSeries:
    category: str
    code: str
    file: str
    kind: str  # index | yield | rate
    duration_years: float | None = None
    duration_note: str | None = None


BCRP_SERIES = (
    BcrpSeries("stocks", "PN01142MM", "acciones__bcrp_indice_general_bvl.csv", "index"),
    # Debt: PN01113MM (private bonds in S/ up to 3 years) and PN01124MM (Treasury bonds up to
    # 5 years) are issuance yields with 0.0 in every month without issues (146/200 and 118/188
    # months in 2010-2026), too sparse for monthly returns. The yield of the outstanding stock
    # of BCRP certificates of deposit is published every month and is used instead.
    BcrpSeries(
        "debt", "PN06503OM", "deuda__bcrp_tasa_saldo_cd_bcrp.csv", "yield", 0.5,
        "Aproximación: duración supuesta de ~0,5 años para el saldo de CD BCRP (plazos mayormente "
        "menores a un año); el fondo de deuda real puede tener mayor duración.",
    ),
    BcrpSeries(
        "bonds", "PD31895MM", "btp__bcrp_rendimiento_10a.csv", "yield", 7.0,
        "Aproximación: duración modificada supuesta de ~7 años para un bono soberano a 10 años.",
    ),
    BcrpSeries("term", "PN07814NM", "deposito__bcrp_tasa_pasiva_181_360d.csv", "rate"),
)
YAHOO_STOCKS_FILE = "acciones__yahoo_epu_en_soles.csv"
# Legacy fallback (W14), used only when the SBS series is unavailable and the manifest has no
# previous mixed entry: 0.5 x BVL index + 0.5 x BTP 10y (duration 7).
MIXED_COMPOSITE = {"components": ["stocks", "bonds"], "weights": [0.5, 0.5]}

# SBS (W15). Month folders use the SBS spelling ("Setiembre"; "Septiembre" returns 404).
SBS_URL = "https://intranet2.sbs.gob.pe/estadistica/financiera/{year}/{folder}/B-220932-{code}{year}.XLS"
SBS_MONTHS = (
    ("Enero", "en"), ("Febrero", "fe"), ("Marzo", "ma"), ("Abril", "ab"), ("Mayo", "my"), ("Junio", "jn"),
    ("Julio", "jl"), ("Agosto", "ag"), ("Setiembre", "se"), ("Octubre", "oc"), ("Noviembre", "no"),
    ("Diciembre", "di"),
)
SBS_MONTHS_BACK = 3  # current month's file, then up to 3 previous months
SBS_USER_AGENT = (
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0 Safari/537.36"
)
SBS_FUND_TYPE = 2
SBS_AFPS = ("Integra", "Prima", "Profuturo")  # averaged; Habitat starts 2013-06 (kept for transparency)
SBS_ALL_AFPS = ("Integra", "Prima", "Profuturo", "Habitat")
SBS_MAX_DAILY_JUMP = 0.5  # a single-day valor cuota change above 50 % is a data error, not a return
SBS_LEVELS_FILE = "mixtos__sbs_afp_fondo2_valor_cuota.csv"
SBS_INDEX_FILE = "mixtos__sbs_afp_fondo2_promedio.csv"
SBS_TITLE = "Valor cuota del Fondo de Pensiones Tipo 2 (promedio Integra, Prima, Profuturo)"
SBS_NOTES = [
    "Aproximación: Fondo 2 de las AFP (fondo mixto por regulación), no un fondo mutuo minorista.",
    "Cerca de 40-50 % del fondo está en activos del exterior: el retorno en soles incluye el efecto del tipo de cambio.",
    "La comisión de la AFP se cobra fuera del valor cuota (no verificado), así que el retorno no la descuenta.",
    "El valor cuota es un precio por cuota: aportes y retiros no distorsionan el retorno.",
    "Promedio simple de los retornos mensuales de Integra, Prima y Profuturo, encadenado como índice "
    "(base 100 en el primer mes); Habitat se excluye porque empieza en 2013-06.",
    "Filtro de datos: se descarta un valor diario que salta más de 50 % frente al anterior y vuelve al "
    "día siguiente; un salto persistente detiene la descarga.",
]


# --------------------------------------------------------------------------- parsing


def parse_spanish_month(label: str) -> str:
    """``'Ene24'`` -> ``'2024-01'``; two-digit years above next year are read as 19xx."""
    text = label.strip()
    month = SPANISH_MONTHS.get(text[:3].lower())
    year_text = text[3:].strip()
    if month is None or not year_text.isdigit() or len(year_text) not in (2, 4):
        raise ValueError(f"unrecognized BCRP period {label!r}")
    year = int(year_text)
    if len(year_text) == 2:
        pivot = (dt.date.today().year + 1) % 100
        year += 2000 if year <= pivot else 1900
    return f"{year:04d}-{month:02d}"


def parse_value(text: str) -> float | None:
    """Published number or ``None`` for "n.d."/empty; accepts ``1,234.5`` and ``6,7``."""
    cleaned = text.strip().replace("\xa0", "").replace(" ", "")
    if cleaned.lower() in MISSING_MARKERS:
        return None
    if "," in cleaned:
        cleaned = cleaned.replace(",", "") if "." in cleaned else cleaned.replace(",", ".")
    try:
        return float(cleaned)
    except ValueError:
        return None


class _BcrpTableParser(HTMLParser):
    """Collect the text of ``<td class="periodo">`` / ``<td class="dato">`` cells in order."""

    def __init__(self) -> None:
        super().__init__()
        self.cells: list[tuple[str, str]] = []
        self.title: str | None = None
        self._cell: str | None = None
        self._text: list[str] = []
        self._in_title = False
        self._in_header = False

    def handle_starttag(self, tag, attrs):
        if tag == "th" and dict(attrs).get("abbr") == "series":
            self._in_header, self._text = True, []
        elif tag == "td":
            classes = (dict(attrs).get("class") or "").split()
            self._cell = "periodo" if "periodo" in classes else "dato" if "dato" in classes else None
            self._text = []
        elif tag == "title":
            self._in_title, self._text = True, []

    def handle_endtag(self, tag):
        if tag == "td" and self._cell is not None:
            self.cells.append((self._cell, "".join(self._text).strip()))
            self._cell = None
        elif tag == "title" and self._in_title:
            self.title = self.title or " ".join("".join(self._text).split())
            self._in_title = False
        elif tag == "th" and self._in_header:
            # The column header carries the full published name (group - series); prefer it.
            self.title, self._in_header = " ".join("".join(self._text).split()), False

    def handle_data(self, data):
        if self._cell is not None or self._in_title or self._in_header:
            self._text.append(data)


def parse_bcrp_table(html: str) -> tuple[str | None, list[tuple[str, float | None]]]:
    """Return the page title and ``[(YYYY-MM, value or None), ...]`` from a BCRP HTML view."""
    parser = _BcrpTableParser()
    parser.feed(html)
    rows: list[tuple[str, float | None]] = []
    period: str | None = None
    for kind, text in parser.cells:
        if kind == "periodo":
            period = text
        elif kind == "dato" and period is not None:
            rows.append((parse_spanish_month(period), parse_value(text)))
            period = None
    return parser.title, rows


# --------------------------------------------------------------------------- SBS parsing


def _cell_text(value) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def parse_sbs_workbook(source, fund_type: int = SBS_FUND_TYPE) -> dict[str, list[tuple[dt.date, float]]]:
    """Daily valor cuota per AFP for one fund type from the B-220932 workbook.

    Layout: a header row labels column groups "Fondo Tipo 0..3"; the next row names
    the AFP of each column and the first column holds the date. Empty cells and
    non-date rows (titles, notes) are skipped. ``source`` is a path or binary file.
    """
    from openpyxl import load_workbook  # local import: data tooling only

    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        label = f"fondo tipo {fund_type}"
        columns: dict[str, int] | None = None
        group_row: tuple | None = None
        series: dict[str, list[tuple[dt.date, float]]] = {}
        for row in rows:
            if columns is None:
                if group_row is not None:
                    start = next(i for i, v in enumerate(group_row) if _cell_text(v).lower() == label)
                    end = next(
                        (i for i, v in enumerate(group_row) if i > start and _cell_text(v).lower().startswith("fondo tipo")),
                        len(row),
                    )
                    columns = {_cell_text(row[i]): i for i in range(start, min(end, len(row))) if _cell_text(row[i])}
                    series = {afp: [] for afp in columns}
                elif any(_cell_text(v).lower() == label for v in row):
                    group_row = row
                continue
            stamp = row[0] if row else None
            if not isinstance(stamp, (dt.datetime, dt.date)):
                continue
            day = stamp.date() if isinstance(stamp, dt.datetime) else stamp
            for afp, index in columns.items():
                value = row[index] if index < len(row) else None
                if isinstance(value, (int, float)) and not isinstance(value, bool) and value > 0:
                    series[afp].append((day, float(value)))
    finally:
        workbook.close()
    if columns is None:
        raise ValueError(f"no 'Fondo Tipo {fund_type}' header found in the SBS workbook")
    return {afp: sorted(points) for afp, points in series.items()}


def reject_daily_spikes(
    points: list[tuple[dt.date, float]], max_jump: float = SBS_MAX_DAILY_JUMP
) -> tuple[list[tuple[dt.date, float]], int]:
    """Drop single-day spikes: a value more than ``max_jump`` away from the last kept one whose
    next value comes back within ``max_jump`` (or that is the last value). A jump that persists
    is a level break, not a typo, and raises ``ValueError`` so a human checks it."""
    kept: list[tuple[dt.date, float]] = []
    rejected = 0
    for i, (day, value) in enumerate(points):
        if kept and abs(value / kept[-1][1] - 1.0) > max_jump:
            following = points[i + 1][1] if i + 1 < len(points) else None
            if following is None or abs(following / kept[-1][1] - 1.0) <= max_jump:
                rejected += 1
                continue
            raise ValueError(f"persistent valor cuota jump on {day}: {kept[-1][1]} -> {value}")
        kept.append((day, value))
    return kept, rejected


def _last_weekday(day: dt.date) -> dt.date:
    """Last Monday-Friday of ``day``'s month."""
    following = dt.date(day.year + day.month // 12, day.month % 12 + 1, 1)
    last = following - dt.timedelta(days=1)
    while last.weekday() >= 5:
        last -= dt.timedelta(days=1)
    return last


def month_end_levels(points: list[tuple[dt.date, float]], current: str, start: str) -> dict[str, float]:
    """Last available value of each month in ``[start, current)`` (``YYYY-MM``).

    The open month is dropped, and so is the series' final month when its data stops
    before that month's last weekday (the workbook was cut before the month closed).
    """
    levels: dict[str, float] = {}
    last_day: dt.date | None = None
    for day, value in sorted(points):
        month = day.strftime("%Y-%m")
        if start <= month < current:
            levels[month] = value
            last_day = day
    if last_day is not None and last_day < _last_weekday(last_day):
        levels.pop(last_day.strftime("%Y-%m"))
    return levels


def _next_month(month: str) -> str:
    year, mon = int(month[:4]), int(month[5:7])
    return f"{year + mon // 12:04d}-{mon % 12 + 1:02d}"


def equal_weight_index(levels: dict[str, dict[str, float]], afps: tuple[str, ...], base: float = 100.0) -> list[tuple[str, float]]:
    """Chain the equal-weight average of the AFPs' monthly returns into an index (``base`` at the first month).

    Uses the months where every AFP has a value; they must be consecutive (a gap would
    hide a return), otherwise ``ValueError``.
    """
    months = sorted(set.intersection(*(set(levels[a]) for a in afps)))
    if len(months) < 2:
        raise ValueError("at least two common months are required")
    index = [(months[0], base)]
    for previous, month in zip(months, months[1:]):
        if month != _next_month(previous):
            raise ValueError(f"gap in the SBS series between {previous} and {month}")
        mean_return = sum(levels[a][month] / levels[a][previous] - 1.0 for a in afps) / len(afps)
        index.append((month, index[-1][1] * (1.0 + mean_return)))
    return index


# --------------------------------------------------------------------------- network


def find_chromium() -> str:
    for name in CHROMIUM_CANDIDATES:
        path = shutil.which(name)
        if path:
            return path
    raise SystemExit(
        "Chromium (or Google Chrome) is required to download BCRP series: the BCRP site only "
        "serves data after a browser challenge. Install it (e.g. 'sudo pacman -S chromium' or "
        "'sudo apt install chromium') or use --stocks-source yahoo for the stocks fallback."
    )


def dump_dom(chromium: str, url: str) -> str:
    command = [
        chromium, "--headless=new", "--no-sandbox", "--disable-gpu",
        "--virtual-time-budget=15000", "--dump-dom", url,
    ]
    result = subprocess.run(command, capture_output=True, text=True, timeout=120, check=False)
    if result.returncode != 0:
        raise RuntimeError(f"chromium exited with {result.returncode} for {url}")
    return result.stdout


def fetch_bcrp(chromium: str, code: str, start: str, end: str) -> tuple[str, str, list[tuple[str, float]]]:
    url = BCRP_URL.format(code=code, start=start, end=end)
    last_error: Exception | None = None
    for _ in range(MAX_ATTEMPTS):
        try:
            title, rows = parse_bcrp_table(dump_dom(chromium, url))
        except (RuntimeError, ValueError, subprocess.TimeoutExpired) as exc:
            last_error = exc
            continue
        if rows:
            return url, title or code, [(m, v) for m, v in rows if v is not None]
        last_error = RuntimeError(f"no data table in the page for {code} (anti-bot challenge not passed?)")
    raise RuntimeError(f"could not download {code}: {last_error}")


def _yahoo_monthly_close(ticker: str) -> dict[str, float]:
    import numpy as np
    import yfinance as yf  # local import: optional dependency

    frame = yf.Ticker(ticker).history(period="max", interval="1mo", auto_adjust=True)
    if frame.empty:
        raise RuntimeError(f"no data returned for {ticker}")
    return {
        stamp.strftime("%Y-%m"): float(value)
        for stamp, value in frame["Close"].items()
        if np.isfinite(value) and value > 0
    }


def fetch_yahoo_stocks() -> list[tuple[str, float]]:
    epu, pen = _yahoo_monthly_close("EPU"), _yahoo_monthly_close("PEN=X")
    return [(m, epu[m] * pen[m]) for m in sorted(set(epu) & set(pen))]


def sbs_candidate_urls(today: dt.date, months_back: int = SBS_MONTHS_BACK) -> list[str]:
    """The current month's B-220932 URL, then up to ``months_back`` previous months."""
    urls = []
    year, month = today.year, today.month
    for _ in range(months_back + 1):
        folder, code = SBS_MONTHS[month - 1]
        urls.append(SBS_URL.format(year=year, folder=folder, code=code))
        year, month = (year - 1, 12) if month == 1 else (year, month - 1)
    return urls


def download_sbs(urls: list[str]) -> tuple[str, bytes]:
    """First URL that returns an .xlsx (zip) payload; one request per URL, no retries (Imperva)."""
    failures = []
    for url in urls:
        request = urllib.request.Request(url, headers={"User-Agent": SBS_USER_AGENT, "Accept": "*/*"})
        try:
            with urllib.request.urlopen(request, timeout=60) as response:
                payload = response.read()
        except (urllib.error.URLError, TimeoutError) as exc:
            failures.append(f"{url}: {exc}")
            continue
        if payload[:2] == b"PK":
            return url, payload
        failures.append(f"{url}: not a workbook ({len(payload)} bytes; blocked by the SBS firewall?)")
    raise RuntimeError("could not download the SBS valor cuota workbook:\n  " + "\n  ".join(failures))


# --------------------------------------------------------------------------- output


def write_csv(path: Path, rows: list[tuple[str, float]]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["date", "value"])
        for month, value in rows:
            writer.writerow([f"{month}-01", repr(value) if isinstance(value, float) else value])


def _span(rows: list[tuple[str, float]]) -> dict:
    return {"first": rows[0][0], "last": rows[-1][0], "months": len(rows)}


def _complete(rows: list[tuple[str, float]], current: str) -> list[tuple[str, float]]:
    return sorted((m, v) for m, v in rows if m < current)


def write_levels_csv(path: Path, levels: dict[str, dict[str, float]], afps: tuple[str, ...]) -> None:
    """Month-end raw levels, one column per AFP (lower-case), empty where an AFP has no value."""
    months = sorted(set().union(*(levels[a].keys() for a in afps)))
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["date", *(a.lower() for a in afps)])
        for month in months:
            writer.writerow([f"{month}-01", *(repr(levels[a][month]) if month in levels[a] else "" for a in afps)])


def _bcrp_start_month(start: str) -> str:
    year, month = start.split("-")
    return f"{int(year):04d}-{int(month):02d}"


def fetch_bcrp_entries(out_dir: Path, start: str, end: str, current: str, fetched_at: str, stocks_source: str) -> dict[str, dict]:
    entries: dict[str, dict] = {}
    chromium = find_chromium()
    for spec in BCRP_SERIES:
        if spec.category == "stocks" and stocks_source == "yahoo":
            continue
        url, title, rows = fetch_bcrp(chromium, spec.code, start, end)
        rows = _complete(rows, current)
        if spec.kind != "index":
            # BCRP writes 0.0 for months without operations in rate series: treat as missing.
            rows = [(m, v) for m, v in rows if v != 0.0]
        if len(rows) < MIN_MONTHS:
            raise RuntimeError(f"{spec.code}: only {len(rows)} months with data")
        write_csv(out_dir / spec.file, rows)
        entry = {"file": spec.file, "source": "BCRP", "series_code": spec.code, "title": title, "kind": spec.kind}
        if spec.duration_years is not None:
            entry["duration_years"] = spec.duration_years
            entry["duration_note"] = spec.duration_note
        entries[spec.category] = {**entry, "url": url, "fetched_at": fetched_at, **_span(rows)}
        print(f"{spec.category:7s} {spec.code}: {len(rows)} months ({rows[0][0]} .. {rows[-1][0]}) -> {spec.file}")

    if stocks_source == "yahoo":
        rows = _complete(fetch_yahoo_stocks(), current)
        write_csv(out_dir / YAHOO_STOCKS_FILE, rows)
        entries["stocks"] = {
            "file": YAHOO_STOCKS_FILE, "source": "Yahoo", "series_code": "EPU*PEN=X",
            "title": "iShares MSCI Peru ETF (EPU) convertido a soles con PEN=X", "kind": "index",
            "url": "https://finance.yahoo.com/quote/EPU", "fetched_at": fetched_at, **_span(rows),
        }
        print(f"stocks  Yahoo EPU*PEN=X: {len(rows)} months -> {YAHOO_STOCKS_FILE}")
    return entries


def build_sbs_series(
    daily: dict[str, list[tuple[dt.date, float]]], current: str, start: str
) -> tuple[dict[str, dict[str, float]], list[tuple[str, float]], dict[str, int]]:
    """Month-end levels per AFP (after the spike filter), the equal-weight index and rejected counts."""
    missing = [a for a in SBS_ALL_AFPS if a not in daily]
    if missing:
        raise ValueError(f"AFP columns missing in the SBS workbook: {missing}")
    levels: dict[str, dict[str, float]] = {}
    rejected: dict[str, int] = {}
    for afp in SBS_ALL_AFPS:
        kept, rejected[afp] = reject_daily_spikes(daily[afp])
        levels[afp] = month_end_levels(kept, current, start)
    return levels, equal_weight_index(levels, SBS_AFPS), rejected


def fetch_sbs_entry(out_dir: Path, start: str, today: dt.date) -> dict:
    current = today.strftime("%Y-%m")
    url, payload = download_sbs(sbs_candidate_urls(today))
    daily = parse_sbs_workbook(io.BytesIO(payload))
    levels, index, rejected = build_sbs_series(daily, current, start)
    if len(index) < MIN_MONTHS:
        raise RuntimeError(f"SBS Fondo 2: only {len(index)} months with data")
    write_levels_csv(out_dir / SBS_LEVELS_FILE, levels, SBS_ALL_AFPS)
    write_csv(out_dir / SBS_INDEX_FILE, index)
    print(f"mixed   SBS B-220932 Fondo 2: {len(index)} months ({index[0][0]} .. {index[-1][0]}) -> {SBS_INDEX_FILE}"
          f" (daily spikes rejected: {rejected})")
    return {
        "file": SBS_INDEX_FILE, "source": "SBS", "series_code": "B-220932 Fondo Tipo 2",
        "title": SBS_TITLE, "kind": "index",
        "url": url, "url_pattern": SBS_URL, "fetched_at": today.isoformat(), **_span(index),
        "afps": list(SBS_AFPS),
        "levels_file": SBS_LEVELS_FILE,
        "excluded_afps": {"Habitat": f"empieza en {min(levels['Habitat'])}; se incluye en {SBS_LEVELS_FILE} solo como referencia"},
        "daily_spikes_rejected": rejected,
        "notes": SBS_NOTES,
    }


def composite_mixed_entry(out_dir: Path, entries: dict[str, dict], fetched_at: str) -> dict:
    """Legacy W14 composite (stocks + bonds), only as a fallback when the SBS series is unavailable."""
    def months(entry: dict) -> set[str]:
        with open(out_dir / entry["file"], newline="", encoding="utf-8") as handle:
            return {row["date"][:7] for row in csv.DictReader(handle) if row["value"]}

    common = sorted(set.intersection(*(months(entries[c]) for c in MIXED_COMPOSITE["components"])))
    return {
        "file": None, "source": "composite", "series_code": None,
        "title": "Compuesto a partir de series reales: 50 % acciones (Índice General BVL) + 50 % bonos (BTP 10 años) (retornos mensuales)",
        "kind": "composite", **MIXED_COMPOSITE,
        "url": None, "fetched_at": fetched_at,
        "first": common[0], "last": common[-1], "months": len(common),
    }


def fetch(out_dir: Path, start: str, stocks_source: str, only: str = "all", today: dt.date | None = None) -> Path:
    today = today or dt.date.today()
    current, end = today.strftime("%Y-%m"), f"{today.year}-{today.month}"
    fetched_at = today.isoformat()
    out_dir.mkdir(parents=True, exist_ok=True)
    target = out_dir / "manifest.json"
    previous = json.loads(target.read_text(encoding="utf-8")).get("categories", {}) if target.is_file() else {}
    entries: dict[str, dict] = {}

    if only in ("all", "bcrp"):
        entries.update(fetch_bcrp_entries(out_dir, start, end, current, fetched_at, stocks_source))
    if only in ("all", "sbs"):
        try:
            entries["mixed"] = fetch_sbs_entry(out_dir, _bcrp_start_month(start), today)
        except (RuntimeError, ValueError, ImportError) as exc:
            if only == "sbs":
                raise SystemExit(f"SBS download failed; the committed snapshot is kept.\n{exc}") from exc
            print(f"WARNING: SBS download failed, keeping the previous mixed entry: {exc}")

    categories = {**previous, **entries}
    if "mixed" not in categories:
        categories["mixed"] = composite_mixed_entry(out_dir, categories, fetched_at)
    missing = [c for c in CATEGORIES if c not in categories]
    if missing:
        raise SystemExit(f"manifest would lack {missing}; run the full download first")
    manifest = {
        "description": "Monthly market series per category (W14, mixed funds from the SBS in W15). Values in "
        "each CSV are exactly as published: index level (kind=index) or % annual rate (kind=yield|rate); "
        "the mixed index is the chained equal-weight average of AFP Fondo 2 valor cuota returns. "
        "ai/uncertainty/bayesian/market.py converts them to monthly returns by kind.",
        "generated_by": "scripts/fetch_market_data.py",
        "categories": {c: categories[c] for c in CATEGORIES},
    }
    target.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"wrote {target}")
    return target


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", type=Path, default=ROOT / "data" / "market")
    parser.add_argument("--start", default="2010-1", help="first month, BCRP format YYYY-M")
    parser.add_argument("--stocks-source", choices=("bcrp", "yahoo"), default="bcrp")
    parser.add_argument(
        "--only", choices=("all", "bcrp", "sbs"), default="all",
        help="download only one source and keep the other manifest entries (default: all)",
    )
    args = parser.parse_args()
    fetch(args.out, args.start, args.stocks_source, args.only)


if __name__ == "__main__":
    main()
